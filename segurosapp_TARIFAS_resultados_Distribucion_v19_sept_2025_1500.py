#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import re
import sys
import threading
import unicodedata
from dataclasses import dataclass
from datetime import datetime, date
from pathlib import Path
from typing import List, Optional, Dict

import numpy as np
import pandas as pd

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from tkinter.scrolledtext import ScrolledText

# ───────────────────────────── Config ─────────────────────────────
APP_TITLE = "Producción – Endosos (solo Producción)"
VERSION = "v1.1"

# ───────────────────── Utilidades de texto/columnas ─────────────────────
def _norm(s: str) -> str:
    if s is None:
        return ""
    s = "".join(c for c in unicodedata.normalize("NFKD", str(s)) if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()

def _normalize_text(x) -> str:
    return _norm(x)

def _find_col_fast(df: pd.DataFrame, *cands: str) -> Optional[str]:
    if df is None or df.empty:
        return None
    norm_map = {_normalize_text(c): c for c in df.columns}
    # exacto
    for cand in cands:
        k = _normalize_text(cand)
        if k in norm_map:
            return norm_map[k]
    # contiene
    for cand in cands:
        k = _normalize_text(cand)
        for key, real in norm_map.items():
            if k and k in key:
                return real
    return None

# ────────────────────────── Lectura robusta ──────────────────────────
from openpyxl import load_workbook
from openpyxl.styles import NamedStyle
from openpyxl.descriptors import String as _StringDesc

_StringDesc.allow_none = True  # evitar errores por estilos corruptos

def _strip_named_styles(wb):
    valid = [st for st in getattr(wb, "_named_styles", []) if isinstance(getattr(st, "name", None), str)]
    if "Normal" not in [st.name for st in valid]:
        valid.insert(0, NamedStyle(name="Normal"))
    wb._named_styles = valid
    wb._named_style_map = {st.name: i for i, st in enumerate(valid)}
    if hasattr(wb, "named_styles"):
        wb.named_styles = valid

def open_and_clean_workbook(path: Path):
    wb = load_workbook(str(path), read_only=True, data_only=True, rich_text=False, keep_links=False)
    _strip_named_styles(wb)
    return wb

_HEADER_TOKENS = {
    "poliza","póliza","endoso","art","artículo","articulo",
    "f/emision","f emision","f emisión","fecha emisión",
    "fec. desde art","fec. hasta art",
    "nombre sección","nombre seccion","sección principal","seccion principal"
}

def _scan_for_header_row(values_2d, max_rows: int = 80) -> int:
    best_row, best_score = 0, -1
    upto = min(max_rows, len(values_2d))
    for r in range(upto):
        row = values_2d[r]
        toks = [_norm(x) for x in row if x not in (None, "")]
        if not toks:
            continue
        kw_hits = sum(any(k in t for k in _HEADER_TOKENS) for t in toks)
        score = kw_hits * 3 + len(set(toks))
        if score > best_score:
            best_score, best_row = score, r
            if kw_hits >= 6:
                break
    return best_row

def _sheet_to_dataframe(ws) -> pd.DataFrame:
    values = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
    if not values:
        return pd.DataFrame()
    hdr_idx = _scan_for_header_row(values, max_rows=80)
    header = values[hdr_idx]
    data = values[hdr_idx+1:]
    data = [row for row in data if any(x not in (None,""," ") for x in row)]
    cols = [str(c).strip() if c is not None else f"Col_{i}" for i, c in enumerate(header)]
    return pd.DataFrame(data, columns=cols)

def robust_read_excel_all_sheets(path: Path) -> pd.DataFrame:
    try:
        wb = open_and_clean_workbook(path)
        frames = []
        for sh in wb.sheetnames:
            try:
                df = _sheet_to_dataframe(wb[sh])
                if not df.empty:
                    frames.append(df)
            except Exception:
                continue
        wb.close()
        if frames:
            return pd.concat(frames, ignore_index=True, sort=False)
    except Exception:
        pass
    # fallback pandas
    try:
        xls = pd.ExcelFile(path, engine="openpyxl")
        frames = []
        for sh in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sh, header=0)
            if not df.empty:
                frames.append(df)
        if frames:
            return pd.concat(frames, ignore_index=True, sort=False)
    except Exception:
        pass
    return pd.DataFrame()

def safe_read_excel(path: Path) -> pd.DataFrame:
    try:
        wb = open_and_clean_workbook(path)
        df = _sheet_to_dataframe(wb.active)
        wb.close()
        if not df.empty:
            return df
        return robust_read_excel_all_sheets(path)
    except Exception:
        return robust_read_excel_all_sheets(path)

def safe_read_csv(path: Path) -> pd.DataFrame:
    for enc in ("utf-8-sig","latin-1","cp1252"):
        try:
            return pd.read_csv(path, sep=None, engine="python", encoding=enc)
        except Exception:
            continue
    return pd.DataFrame()

# ────────────────────── Núcleo de negocio: Producción ──────────────────────
def _to_num(series: pd.Series) -> pd.Series:
    s = series.astype(str)
    s = s.str.replace(r"[^\d\-\+eE\.,]", "", regex=True)
    s = s.str.replace(r"\.(?=.*\,)", "", regex=True)  # miles con punto si hay coma decimal
    s = s.str.replace(",", ".", regex=False)
    return pd.to_numeric(s, errors="coerce")

def _to_date(series: pd.Series) -> pd.Series:
    return pd.to_datetime(series, dayfirst=True, errors="coerce").dt.date

def _ensure_string_headers_unique(df: pd.DataFrame) -> pd.DataFrame:
    cols, seen = [], {}
    for c in df.columns:
        k = str(c) if c is not None else "Col"
        if k in seen:
            seen[k] += 1
            k = f"{k}_{seen[k]}"
        else:
            seen[k] = 0
        cols.append(k)
    df = df.copy(); df.columns = cols
    return df

def export_excel_clean(path: Path, sheets: Dict[str, pd.DataFrame]) -> None:
    """
    Exporta usando xlsxwriter si está disponible (rápido).
    Si no, cae a openpyxl (compatible).
    """
    # Intento 1: xlsxwriter (pandas >= cualquier versión; sin 'options')
    try:
        with pd.ExcelWriter(path, engine="xlsxwriter") as writer:
            for name, df in sheets.items():
                if df is None:
                    continue
                _ensure_string_headers_unique(df).to_excel(
                    writer, sheet_name=str(name)[:31], index=False
                )
        return
    except Exception:
        pass

    # Intento 2: openpyxl puro
    from openpyxl import Workbook
    wb = Workbook()
    ws0 = wb.active
    wb.remove(ws0)
    for name, df in sheets.items():
        ws = wb.create_sheet(title=str(name)[:31])
        if df is None or df.empty:
            continue
        df2 = _ensure_string_headers_unique(df)
        ws.append(list(df2.columns))
        for row in df2.itertuples(index=False, name=None):
            ws.append(list(row))
    wb.save(str(path))
    wb.close()

def _depurada_operacion_algebra(df_base: pd.DataFrame) -> pd.DataFrame:
    dep = df_base.copy()
    if "Artículo" not in dep.columns and "Art." in dep.columns:
        dep.rename(columns={"Art.": "Artículo"}, inplace=True)

    # Eliminar Tipo Póliza = 5 ANTES de agrupar
    tcol_code = _find_col_fast(dep, "Tipo Póliza", "Tipo Poliza", "Tipo")
    if tcol_code:
        dep = dep[dep[tcol_code].astype(str).str.strip() != "5"].copy()

    # Señalización anulaciones y modificaciones
    tname = _find_col_fast(dep, "Nombre Tipo Póliza", "Nombre Tipo Poliza")
    tnorm = dep[tname].map(_normalize_text) if tname else pd.Series("", index=dep.index)
    is_mod = tnorm.str.contains(r"\bendoso\b.*\bmodificaci", na=False) | tnorm.str.fullmatch(r".*\bmodificaci.*", na=False)
    dep = dep[~is_mod].copy()
    tnorm = tnorm.loc[dep.index]

    is_anul = tnorm.str.contains(r"\banulaci[oó]n\b", na=False)
    if tcol_code:
        is_anul = is_anul | dep[tcol_code].astype(str).str.strip().eq("3")

    # Números
    for colname in ["Prima Art.", "Costo de Servicio Art.", "Importe Agente Art."]:
        col = _find_col_fast(dep, colname)
        if col:
            dep[col] = _to_num(dep[col])

    # Fechas como date
    col_fd = _find_col_fast(dep, "Fec. Desde Art.", "Fecha Desde Art.", "Inicio Vigencia")
    col_fh = _find_col_fast(dep, "Fec. Hasta Art.", "Fecha Hasta Art.", "Fin Vigencia")
    col_emi = _find_col_fast(dep, "F/Emisión", "F Emisión", "F Emision", "Fecha Emisión", "Fecha emision")
    if col_fd: dep[col_fd] = _to_date(dep[col_fd])
    if col_fh: dep[col_fh] = _to_date(dep[col_fh])
    if col_emi: dep[col_emi] = _to_date(dep[col_emi])

    # Agrupación clave
    keys = [_find_col_fast(dep, "Nombre Sección"), _find_col_fast(dep, "Póliza"), _find_col_fast(dep, "Artículo")]
    keys = [k for k in keys if k]
    g = dep.groupby(keys, dropna=False) if keys else None

    # Sumas por artículo
    col_prima = _find_col_fast(dep, "Prima Art.")
    if col_prima and g is not None:
        dep["Suma Art."] = g[col_prima].transform("sum")

    col_costo = _find_col_fast(dep, "Costo de Servicio Art.")
    if col_costo and g is not None:
        dep["Suma Costo de Servicio Art."] = g[col_costo].transform("sum")

    col_imp = _find_col_fast(dep, "Importe Agente Art.")
    if col_imp and g is not None:
        dep["Suma Importe Agente Art."] = g[col_imp].transform("sum")

    # Fechas inicio/fin
    fini_ts = pd.to_datetime(dep[col_fd], errors="coerce") if col_fd else pd.Series(pd.NaT, index=dep.index)
    fh_ts   = pd.to_datetime(dep[col_fh], errors="coerce") if col_fh else pd.Series(pd.NaT, index=dep.index)

    dep["_fIni_ts"] = fini_ts
    dep["_fIniNoAnul_ts"] = dep["_fIni_ts"].where(~is_anul)
    if g is not None:
        dep["fechaInicio"] = g["_fIniNoAnul_ts"].transform("min").dt.date
    else:
        dep["fechaInicio"] = dep["_fIniNoAnul_ts"].dt.date

    dep["_fDesdeAnu_ts"] = dep["_fIni_ts"].where(is_anul)  # fecha desde de las anulaciones
    if g is not None:
        dep["_fDesdeAnuMin_ts"] = g["_fDesdeAnu_ts"].transform("min")
    else:
        dep["_fDesdeAnuMin_ts"] = dep["_fDesdeAnu_ts"]

    dep["fechaFin"] = dep["_fDesdeAnuMin_ts"].dt.date  # si hubo anulaciones
    if col_fh:
        dep["fechaFin"] = dep["fechaFin"].where(dep["fechaFin"].notna(), fh_ts.dt.date)

    # Quitar anulaciones del resultado final
    dep = dep[~is_anul].copy()

    dep.drop(columns=[c for c in ["_fIni_ts","_fIniNoAnul_ts","_fDesdeAnu_ts","_fDesdeAnuMin_ts"]
                      if c in dep.columns], inplace=True, errors="ignore")

    # Orden recomendado
    col_pol = _find_col_fast(dep, "Póliza","Poliza")
    col_art = _find_col_fast(dep, "Artículo","Art.")
    try:
        if col_pol and col_art:
            dep = dep.sort_values(by=[col_pol, col_art], ignore_index=True)
    except Exception:
        pass

    # Mantener columnas originales + calculadas
    original_cols = list(df_base.columns)
    calc_cols = [c for c in ["fechaInicio","fechaFin","Suma Art.","Suma Costo de Servicio Art.","Suma Importe Agente Art."] if c in dep.columns]
    final_cols = [c for c in original_cols if c in dep.columns] + [c for c in calc_cols if c not in original_cols]
    return dep[final_cols]

def _exportar_por_categoria(out_dir: Path, prefix: str, ts: str,
                            df_proc: pd.DataFrame, df_orig: pd.DataFrame,
                            log) -> List[Path]:
    gen: List[Path] = []
    campo_dep  = _find_col_fast(df_proc, "Nombre Sección Principal")
    campo_orig = _find_col_fast(df_orig, "Nombre Sección Principal")

    if campo_dep is None and campo_orig is None:
        log("⚠️ No existe 'Nombre Sección Principal'; no se generan libros por categoría.")
        return gen

    if campo_dep and not df_proc.empty:
        cats = sorted([c for c in df_proc[campo_dep].dropna().unique()], key=lambda x: str(x))
    elif campo_orig:
        cats = sorted([c for c in df_orig[campo_orig].dropna().unique()], key=lambda x: str(x))
    else:
        cats = []

    if not cats:
        log("⚠️ No hay categorías para exportar.")
        return gen

    log(f"• Exportando {len(cats)} libro(s) por categoría…")
    for cat in cats:
        safe = re.sub(r'[\\/*?:"<>|]', "_", str(cat)).strip()[:80] or "CAT"
        fpath = out_dir / f"{prefix}_{safe}_{ts}.xlsx"
        df_p = df_proc[df_proc[campo_dep] == cat] if campo_dep else pd.DataFrame()
        df_o = df_orig[df_orig[campo_orig] == cat] if campo_orig else pd.DataFrame()

        # Orden en PROCESADO
        col_pol = _find_col_fast(df_p, "Póliza","Poliza")
        col_art = _find_col_fast(df_p, "Artículo","Art.")
        if col_pol and col_art and not df_p.empty:
            df_p = df_p.sort_values(by=[col_pol, col_art], ignore_index=True)

        export_excel_clean(fpath, {"PROCESADO": df_p, "ORIGINAL": df_o})
        gen.append(fpath)
        log(f"  - {safe}: {fpath.name}")
    return gen

def procesar_produccion(files: List[Path], out_dir: Path, prefix: str, log, prog) -> List[Path]:
    generated_files: List[Path] = []
    raw_dfs: List[pd.DataFrame] = []
    tot = len(files)
    if tot == 0:
        log("No se seleccionaron archivos de producción.")
        return generated_files

    # 1) Leer
    for i, f in enumerate(files, 1):
        prog(i, tot, f"Leyendo {Path(f).name}")
        try:
            suf = Path(f).suffix.lower()
            if suf in (".xlsx",".xlsm",".xls"):
                df = safe_read_excel(Path(f))
            elif suf == ".csv":
                df = safe_read_csv(Path(f))
            else:
                df = safe_read_excel(Path(f))
            if df.empty:
                log(f"{Path(f).name}: archivo vacío/ilegible.")
                continue
            raw_dfs.append(df)
        except Exception as e:
            log(f"{Path(f).name}: error de lectura → {e}")

    if not raw_dfs:
        log("No hubo datos válidos.")
        return generated_files

    # 2) Procesar
    df_orig = pd.concat(raw_dfs, ignore_index=True).drop_duplicates(keep="first").reset_index(drop=True)
    cols_base = [
        "Nombre Sección Principal","Sección","Nombre Sección","Póliza","Endoso",
        "Tipo Póliza","Nombre Tipo Póliza","F/Emisión","Nombre Agente Principal",
        "Artículo","Fec. Desde Art.","Fec. Hasta Art.",
        "Costo de Servicio Art.","Prima Art.","Suma Asegurada Art.","Importe Agente Art."
    ]
    use_cols = [c for c in cols_base if c in df_orig.columns]
    df_base = df_orig[use_cols].copy() if use_cols else df_orig.copy()
    dep = _depurada_operacion_algebra(df_base)

    # Orden final
    col_pol = _find_col_fast(dep, "Póliza","Poliza")
    col_art = _find_col_fast(dep, "Artículo","Art.")
    if col_pol and col_art:
        dep = dep.sort_values(by=[col_pol, col_art], ignore_index=True)

    # 3) Exportar
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    consolidado = out_dir / f"{prefix}_COMPILADO_{ts}.xlsx"
    export_excel_clean(consolidado, {"ORIGINAL": df_orig, "PROCESADO": dep})
    generated_files.append(consolidado)
    log(f"✅ Consolidado: {consolidado.name}")

    # 4) Por categoría
    generated_files += _exportar_por_categoria(out_dir, prefix, ts, dep, df_orig, log)

    prog(1, 1, "Finalizado")
    return generated_files

# ──────────────────────────────── UI ────────────────────────────────
@dataclass
class AppState:
    files: List[Path]
    out_dir: Path | None
    prefix: str

class ProduccionGUI(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(f"{APP_TITLE} {VERSION}")
        self.minsize(900, 520)
        self.geometry("1000x600")
        try:
            self.tk.call("tk", "scaling", 1.2)  # mejor legibilidad en Windows
        except Exception:
            pass

        self.state = AppState(files=[], out_dir=None, prefix="PROD")
        self._build_ui()

    def _build_ui(self):
        s = ttk.Style(self)
        try: s.theme_use("clam")
        except Exception: pass

        root = ttk.Frame(self, padding=12)
        root.pack(fill="both", expand=True)

        for c in (0,1,2):
            root.columnconfigure(c, weight=1 if c == 1 else 0)
        root.rowconfigure(5, weight=1)

        ttk.Label(root, text="Archivos de entrada (Excel/CSV):").grid(row=0, column=0, sticky="w", padx=(0,8), pady=6)
        self.var_files = tk.StringVar(value="0 archivo(s) seleccionado(s)")
        ttk.Entry(root, textvariable=self.var_files, state="readonly").grid(row=0, column=1, sticky="we", pady=6)
        ttk.Button(root, text="Agregar…", command=self._add_files).grid(row=0, column=2, sticky="e", padx=(8,0), pady=6)

        ttk.Label(root, text="Carpeta de salida:").grid(row=1, column=0, sticky="w", padx=(0,8), pady=6)
        self.var_out = tk.StringVar()
        ttk.Entry(root, textvariable=self.var_out, state="readonly").grid(row=1, column=1, sticky="we", pady=6)
        ttk.Button(root, text="Seleccionar…", command=self._choose_out).grid(row=1, column=2, sticky="e", padx=(8,0), pady=6)

        ttk.Label(root, text="Prefijo archivo(s):").grid(row=2, column=0, sticky="w", padx=(0,8), pady=6)
        self.var_prefix = tk.StringVar(value="PROD")
        ttk.Entry(root, textvariable=self.var_prefix, width=20).grid(row=2, column=1, sticky="w", pady=6)

        btns = ttk.Frame(root); btns.grid(row=3, column=0, columnspan=3, sticky="e", pady=(4,6))
        self.btn_run = ttk.Button(btns, text="Procesar", command=self._run, width=14)
        self.btn_run.grid(row=0, column=0, padx=6)
        ttk.Button(btns, text="Limpiar lista", command=self._clear_list, width=14).grid(row=0, column=1, padx=6)
        ttk.Button(btns, text="Salir", command=self.destroy, width=10).grid(row=0, column=2, padx=6)

        self.progress = ttk.Progressbar(root, orient="horizontal", mode="determinate", maximum=100)
        self.progress.grid(row=4, column=0, columnspan=3, sticky="we", pady=(0,6))

        lf = ttk.LabelFrame(root, text="Registro", padding=8)
        lf.grid(row=5, column=0, columnspan=3, sticky="nsew")
        self.log = ScrolledText(lf, height=12, font=("Consolas", 10))
        self.log.pack(fill="both", expand=True)

    # --------------- UI helpers ---------------
    def _add_files(self):
        fl = filedialog.askopenfilenames(title="Seleccionar producción",
                                         filetypes=[("Excel/CSV", "*.xlsx *.xlsm *.xls *.csv")])
        if not fl:
            return
        cur = {p.resolve() for p in self.state.files}
        for p in fl:
            P = Path(p)
            if P.resolve() not in cur:
                self.state.files.append(P)
        self.var_files.set(f"{len(self.state.files)} archivo(s) seleccionado(s)")

    def _choose_out(self):
        d = filedialog.askdirectory(title="Carpeta de salida")
        if d:
            self.state.out_dir = Path(d)
            self.var_out.set(d)

    def _clear_list(self):
        self.state.files = []
        self.var_files.set("0 archivo(s) seleccionado(s)")

    def _log(self, msg: str):
        self.log.insert("end", msg + "\n")
        self.log.see("end")

    def _set_progress(self, val: int):
        self.progress["value"] = max(0, min(100, int(val)))
        self.update_idletasks()

    # --------------- Run ---------------
    def _run(self):
        if not self.state.files:
            messagebox.showerror("Producción", "Agrega uno o más archivos de entrada.")
            return
        if not self.state.out_dir:
            messagebox.showerror("Salida", "Selecciona una carpeta de salida.")
            return
        self.state.prefix = self.var_prefix.get().strip() or "PROD"

        self.btn_run.configure(state="disabled")
        self._set_progress(0)
        self.log.delete("1.0", "end")
        self._log(f"Leyendo {len(self.state.files)} archivo(s)…")

        threading.Thread(target=self._worker, daemon=True).start()

    def _worker(self):
        try:
            gen = procesar_produccion(
                self.state.files,
                self.state.out_dir,
                self.state.prefix,
                log=lambda m: self.after(0, self._log, m),
                prog=lambda cur, tot, _: self.after(0, self._set_progress, int((cur/max(tot,1))*90))
            )
            if gen:
                self.after(0, self._set_progress, 100)
                self.after(0, messagebox.showinfo, "Completado", f"Se generaron {len(gen)} archivo(s) en:\n{self.state.out_dir}")
            else:
                self.after(0, messagebox.showwarning, "Atención", "No se generaron archivos.")
        except Exception as e:
            import traceback
            self.after(0, self._log, f"Error: {e}\n{traceback.format_exc()}")
        finally:
            self.after(0, self.btn_run.configure, {"state": "normal"})

# ──────────────────────────────── Main ────────────────────────────────
if __name__ == "__main__":
    app = ProduccionGUI()
    app.mainloop()
