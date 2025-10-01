#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import re
import sys
import threading
import unicodedata
from dataclasses import dataclass
from datetime import datetime, date
from pathlib import Path
from typing import List, Optional, Dict, Set

import numpy as np
import pandas as pd

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from tkinter.scrolledtext import ScrolledText

# ───────── Excel: lectura robusta y fix estilos ─────────
from openpyxl import load_workbook, Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import NamedStyle
from openpyxl.descriptors import String as _StringDesc
_StringDesc.allow_none = True  # evita errores por estilos corruptos


# ───────────────────────────── Config ─────────────────────────────
APP_TITLE = "Producción – Solo Producción (con selector)"
VERSION = "v1.2"


# ───────────────────── Utilidades de texto/columnas ─────────────────────
def _norm(s: str) -> str:
    if s is None:
        return ""
    s = "".join(c for c in unicodedata.normalize("NFKD", str(s)) if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()


def _normalize_text(x) -> str:
    return _norm(x)


def _find_col_fast(df: pd.DataFrame, *cands: str) -> Optional[str]:
    if df is None or df.empty:
        return None
    norm_map = {_normalize_text(c): c for c in df.columns}
    # exacto
    for cand in cands:
        k = _normalize_text(cand)
        if k in norm_map:
            return norm_map[k]
    # contiene
    for cand in cands:
        k = _normalize_text(cand)
        for key, real in norm_map.items():
            if k and k in key:
                return real
    return None


# ────────────────────────── Lectura robusta ──────────────────────────
def _strip_named_styles(wb):
    valid = [st for st in getattr(wb, "_named_styles", []) if isinstance(getattr(st, "name", None), str)]
    if "Normal" not in [st.name for st in valid]:
        valid.insert(0, NamedStyle(name="Normal"))
    wb._named_styles = valid
    wb._named_style_map = {st.name: i for i, st in enumerate(valid)}
    if hasattr(wb, "named_styles"):
        wb.named_styles = valid


def open_and_clean_workbook(path: Path):
    wb = load_workbook(str(path), read_only=True, data_only=True, rich_text=False, keep_links=False)
    _strip_named_styles(wb)
    return wb


_HEADER_TOKENS = {
    "poliza", "póliza", "endoso", "art", "artículo", "articulo",
    "f/emision", "f emision", "f emisión", "fecha emisión",
    "fec. desde art", "fec. hasta art",
    "nombre sección", "nombre seccion", "sección principal", "seccion principal",
    "nombre sección principal", "nombre seccion principal"
}


def _scan_for_header_row(values_2d, max_rows: int = 80) -> int:
    best_row, best_score = 0, -1
    upto = min(max_rows, len(values_2d))
    for r in range(upto):
        row = values_2d[r]
        toks = [_norm(x) for x in row if x not in (None, "")]
        if not toks:
            continue
        kw_hits = sum(any(k in t for k in _HEADER_TOKENS) for t in toks)
        score = kw_hits * 3 + len(set(toks))
        if score > best_score:
            best_score, best_row = score, r
    return best_row


def _sheet_to_dataframe(ws) -> pd.DataFrame:
    values = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
    if not values:
        return pd.DataFrame()
    hdr_idx = _scan_for_header_row(values, max_rows=80)
    header = values[hdr_idx]
    data = values[hdr_idx + 1:]
    data = [row for row in data if any(x not in (None, "", " ") for x in row)]
    cols = [str(c).strip() if c is not None else f"Col_{i}" for i, c in enumerate(header)]
    return pd.DataFrame(data, columns=cols)


def robust_read_excel_all_sheets(path: Path) -> pd.DataFrame:
    try:
        wb = open_and_clean_workbook(path)
        frames = []
        for sh in wb.sheetnames:
            try:
                df = _sheet_to_dataframe(wb[sh])
                if not df.empty:
                    frames.append(df)
            except Exception:
                continue
        wb.close()
        if frames:
            return pd.concat(frames, ignore_index=True, sort=False)
    except Exception:
        pass
    # fallback pandas
    try:
        xls = pd.ExcelFile(path, engine="openpyxl")
        frames = []
        for sh in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sh, header=0)
            if not df.empty:
                frames.append(df)
        if frames:
            return pd.concat(frames, ignore_index=True, sort=False)
    except Exception:
        pass
    return pd.DataFrame()


def safe_read_excel(path: Path) -> pd.DataFrame:
    try:
        wb = open_and_clean_workbook(path)
        df = _sheet_to_dataframe(wb.active)
        wb.close()
        if not df.empty:
            return df
        return robust_read_excel_all_sheets(path)
    except Exception:
        return robust_read_excel_all_sheets(path)


def safe_read_csv(path: Path) -> pd.DataFrame:
    for enc in ("utf-8-sig", "latin-1", "cp1252"):
        try:
            return pd.read_csv(path, sep=None, engine="python", encoding=enc)
        except Exception:
            continue
    return pd.DataFrame()


# ────────────────────── Núcleo de negocio: Producción ──────────────────────
def _to_num(series: pd.Series) -> pd.Series:
    s = series.astype(str)
    s = s.str.replace(r"[^\d\-\+eE\.,]", "", regex=True)
    s = s.str.replace(r"\.(?=.*\,)", "", regex=True)  # miles con punto si hay coma decimal
    s = s.str.replace(",", ".", regex=False)
    return pd.to_numeric(s, errors="coerce")


def _to_dt(series: pd.Series) -> pd.Series:
    return pd.to_datetime(series, dayfirst=True, errors="coerce")


def _fmt_dt_txt(series: pd.Series) -> pd.Series:
    s = _to_dt(series)
    return s.dt.strftime("%d/%m/%Y")


def _ensure_string_headers_unique(df: pd.DataFrame) -> pd.DataFrame:
    cols, seen = [], {}
    for c in df.columns:
        k = str(c) if c is not None else "Col"
        if k in seen:
            seen[k] += 1
            k = f"{k}_{seen[k]}"
        else:
            seen[k] = 0
        cols.append(k)
    df = df.copy()
    df.columns = cols
    return df


def export_excel_clean(path: Path, sheets: Dict[str, pd.DataFrame]) -> None:
    """
    Exporta usando openpyxl (compatibilidad amplia).
    Evita el uso de pd.ExcelWriter con 'options'.
    """
    wb = Workbook(write_only=False)
    # elimina hoja por defecto
    ws0 = wb.active
    wb.remove(ws0)

    for sheet_name, df in sheets.items():
        ws = wb.create_sheet(title=str(sheet_name)[:31])
        if df is None or df.empty:
            continue
        df2 = _ensure_string_headers_unique(df)
        for r in dataframe_to_rows(df2.where(pd.notna(df2), None), index=False, header=True):
            ws.append(r)

    wb.save(str(path))
    wb.close()


def _depurada_operacion_algebra(df_base: pd.DataFrame) -> pd.DataFrame:
    """
    Lógica de depuración de producción:
    - Renombra Art. -> Artículo si hace falta
    - Elimina Tipo Póliza == 5 (antes de agrupar)
    - Quita modificaciones
    - Señaliza anulaciones (por texto o código == 3)
    - Calcula Suma Art., Suma Costo/Importe por agrupación (Sección/Poliza/Artículo)
    - Calcula fechaInicio = min(Fec. Desde) (sin anulaciones),
      fechaFin = min(Fec. Desde) de anulaciones emparejadas; si no, usa Fec. Hasta original
    - Quita las anulaciones del resultado
    """
    dep = df_base.copy()
    if "Artículo" not in dep.columns and "Art." in dep.columns:
        dep.rename(columns={"Art.": "Artículo"}, inplace=True)

    # Eliminar Tipo Póliza = 5 ANTES de agrupar
    tcol_code = _find_col_fast(dep, "Tipo Póliza", "Tipo Poliza", "Tipo")
    if tcol_code:
        dep = dep[dep[tcol_code].astype(str).str.strip() != "5"].copy()

    # Señalización anulaciones y modificaciones
    tname = _find_col_fast(dep, "Nombre Tipo Póliza", "Nombre Tipo Poliza")
    tnorm = dep[tname].map(_normalize_text) if tname else pd.Series("", index=dep.index)

    is_mod = tnorm.str.contains(r"\bendoso\b.*\bmodificaci", na=False) | tnorm.str.fullmatch(r".*\bmodificaci.*", na=False)
    dep = dep[~is_mod].copy()
    tnorm = tnorm.loc[dep.index]

    is_anul = tnorm.str.contains(r"\banulaci[oó]n\b", na=False)
    if tcol_code:
        is_anul = is_anul | dep[tcol_code].astype(str).str.strip().eq("3")

    # Números
    for colname in ["Prima Art.", "Costo de Servicio Art.", "Importe Agente Art."]:
        col = _find_col_fast(dep, colname)
        if col:
            dep[col] = _to_num(dep[col])

    # Fechas como ts
    col_fd = _find_col_fast(dep, "Fec. Desde Art.", "Fecha Desde Art.", "Inicio Vigencia")
    col_fh = _find_col_fast(dep, "Fec. Hasta Art.", "Fecha Hasta Art.", "Fin Vigencia")
    col_emi = _find_col_fast(dep, "F/Emisión", "F Emisión", "F Emision", "Fecha Emisión", "Fecha emision")
    if col_fd:
        dep[col_fd] = _to_dt(dep[col_fd])
    if col_fh:
        dep[col_fh] = _to_dt(dep[col_fh])
    if col_emi:
        dep[col_emi] = _to_dt(dep[col_emi])

    # Agrupación clave
    keys = [_find_col_fast(dep, "Nombre Sección"), _find_col_fast(dep, "Póliza"), _find_col_fast(dep, "Artículo")]
    keys = [k for k in keys if k]
    g = dep.groupby(keys, dropna=False) if keys else None

    # Sumas por artículo
    col_prima = _find_col_fast(dep, "Prima Art.")
    if col_prima and g is not None:
        dep["Suma Art."] = g[col_prima].transform("sum")

    col_costo = _find_col_fast(dep, "Costo de Servicio Art.")
    if col_costo and g is not None:
        dep["Suma Costo de Servicio Art."] = g[col_costo].transform("sum")

    col_imp = _find_col_fast(dep, "Importe Agente Art.")
    if col_imp and g is not None:
        dep["Suma Importe Agente Art."] = g[col_imp].transform("sum")

    # Fechas inicio/fin
    fini_ts = dep[col_fd] if col_fd else pd.Series(pd.NaT, index=dep.index)
    fh_ts = dep[col_fh] if col_fh else pd.Series(pd.NaT, index=dep.index)

    dep["_fIni_ts"] = fini_ts
    dep["_fIniNoAnul_ts"] = dep["_fIni_ts"].where(~is_anul)
    if g is not None:
        dep["fechaInicio"] = g["_fIniNoAnul_ts"].transform("min").dt.date
    else:
        dep["fechaInicio"] = dep["_fIniNoAnul_ts"].dt.date

    dep["_fDesdeAnu_ts"] = dep["_fIni_ts"].where(is_anul)  # fecha desde de las anulaciones
    if g is not None:
        dep["_fDesdeAnuMin_ts"] = g["_fDesdeAnu_ts"].transform("min")
    else:
        dep["_fDesdeAnuMin_ts"] = dep["_fDesdeAnu_ts"]

    dep["fechaFin"] = dep["_fDesdeAnuMin_ts"].dt.date  # si hubo anulaciones
    if col_fh:
        dep["fechaFin"] = dep["fechaFin"].where(dep["fechaFin"].notna(), fh_ts.dt.date)

    # Quitar anulaciones del resultado final
    dep = dep[~is_anul].copy()

    # Fechas a texto dd/mm/yyyy en columnas originales
    for c in [col_emi, col_fd, col_fh]:
        if c and c in dep.columns:
            dep[c] = _fmt_dt_txt(dep[c])

    dep.drop(columns=[c for c in ["_fIni_ts", "_fIniNoAnul_ts", "_fDesdeAnu_ts", "_fDesdeAnuMin_ts"]
                      if c in dep.columns], inplace=True, errors="ignore")

    # Orden recomendado
    col_pol = _find_col_fast(dep, "Póliza", "Poliza")
    col_art = _find_col_fast(dep, "Artículo", "Art.")
    try:
        if col_pol and col_art:
            dep = dep.sort_values(by=[col_pol, col_art], ignore_index=True)
    except Exception:
        pass

    # Mantener columnas originales + calculadas
    original_cols = list(df_base.columns)
    calc_cols = [c for c in ["fechaInicio", "fechaFin", "Suma Art.", "Suma Costo de Servicio Art.", "Suma Importe Agente Art."]
                 if c in dep.columns]
    final_cols = [c for c in original_cols if c in dep.columns] + [c for c in calc_cols if c not in original_cols]
    return dep[final_cols]


def _pick_cat_column(df: pd.DataFrame, prefer_col: str = "auto") -> Optional[str]:
    """
    prefer_col: 'auto' → prioriza 'Nombre Sección Principal'; 
                'principal' → usa 'Nombre Sección Principal';
                'seccion' → usa 'Nombre Sección'.
    """
    if prefer_col == "principal":
        return _find_col_fast(df, "Nombre Sección Principal", "Nombre Seccion Principal")
    if prefer_col == "seccion":
        return _find_col_fast(df, "Nombre Sección", "Nombre Seccion")
    # auto
    return (_find_col_fast(df, "Nombre Sección Principal", "Nombre Seccion Principal")
            or _find_col_fast(df, "Nombre Sección", "Nombre Seccion"))


def _exportar_por_categoria(out_dir: Path, prefix: str, ts: str,
                            df_proc: pd.DataFrame, df_orig: pd.DataFrame,
                            log, prog,
                            cats_filter: Optional[Set[str]] = None,
                            prefer_col: str = "auto") -> List[Path]:
    gen: List[Path] = []

    campo_dep = _pick_cat_column(df_proc, prefer_col) if df_proc is not None and not df_proc.empty else None
    campo_orig = _pick_cat_column(df_orig, prefer_col) if df_orig is not None and not df_orig.empty else None

    if campo_dep is None and campo_orig is None:
        log("⚠️ No existe 'Nombre Sección Principal' ni 'Nombre Sección'; no se generan libros por categoría.")
        return gen

    cats = set()
    if campo_dep:
        cats |= set(df_proc[campo_dep].dropna().astype(str).unique())
    if campo_orig:
        cats |= set(df_orig[campo_orig].dropna().astype(str).unique())
    cats = sorted(cats, key=lambda x: str(x))

    if cats_filter:
        target = {str(x) for x in cats_filter}
        cats = [c for c in cats if str(c) in target]
        if not cats:
            log("⚠️ Ninguna categoría seleccionada coincide con los datos.")
            return gen

    if not cats:
        log("⚠️ No hay categorías para exportar.")
        return gen

    log(f"• Exportando {len(cats)} libro(s) por categoría…")
    for i, cat in enumerate(cats, 1):
        prog(i, len(cats), f"Exportando categoría: {str(cat)[:40]}")
        safe = re.sub(r'[\\/*?:"<>|]', "_", str(cat)).strip()[:80] or "CAT"
        fpath = out_dir / f"{prefix}_{safe}_{ts}.xlsx"
        df_p = df_proc[df_proc[campo_dep] == cat] if campo_dep else pd.DataFrame()
        df_o = df_orig[df_orig[campo_orig] == cat] if campo_orig else pd.DataFrame()

        # Orden en PROCESADO
        col_pol = _find_col_fast(df_p, "Póliza", "Poliza")
        col_art = _find_col_fast(df_p, "Artículo", "Art.")
        if col_pol and col_art and not df_p.empty:
            df_p = df_p.sort_values(by=[col_pol, col_art], ignore_index=True)

        export_excel_clean(fpath, {"PROCESADO": df_p, "ORIGINAL": df_o})
        gen.append(fpath)
        log(f"  - {safe}: {fpath.name}")
    return gen


def procesar_produccion(files: List[Path] | List[str],
                        out_dir: Path,
                        prefix: str,
                        log,
                        prog,
                        cats_filter: Optional[Set[str]] = None,
                        prefer_col: str = "auto") -> List[Path]:
    generated_files: List[Path] = []
    raw_dfs: List[pd.DataFrame] = []

    tot = len(files)
    if tot == 0:
        log("No se seleccionaron archivos de producción.")
        return generated_files

    for i, f in enumerate(files, 1):
        prog(i, tot, f"Leyendo {Path(f).name}")
        try:
            df = pd.DataFrame()
            if str(f).lower().endswith((".xlsx", ".xlsm", ".xls")):
                df = safe_read_excel(Path(f))
            elif str(f).lower().endswith(".csv"):
                df = safe_read_csv(Path(f))
            else:
                df = safe_read_excel(Path(f))
            if df.empty:
                log(f"{Path(f).name}: archivo vacío/ilegible.")
                continue
            raw_dfs.append(df)
        except Exception as e:
            log(f"{Path(f).name}: error de lectura → {e}")

    if not raw_dfs:
        log("No hubo datos válidos de Producción.")
        return generated_files

    df_orig = pd.concat(raw_dfs, ignore_index=True, sort=False)
    df_orig_clean = df_orig.drop_duplicates(keep="first").reset_index(drop=True)

    # Mantener columnas base si existen (opcional)
    PROD_COLS_BASE = [
        "Nombre Sección Principal", "Sección", "Nombre Sección", "Póliza", "Endoso",
        "Tipo Póliza", "Nombre Tipo Póliza", "F/Emisión", "Nombre Agente Principal",
        "Artículo", "Fec. Desde Art.", "Fec. Hasta Art.",
        "Costo de Servicio Art.", "Prima Art.", "Suma Asegurada Art.", "Importe Agente Art."
    ]
    cols_exist = [c for c in PROD_COLS_BASE if c in df_orig_clean.columns]
    df_base = df_orig_clean[cols_exist].copy() if cols_exist else df_orig_clean.copy()

    if "Artículo" not in df_base.columns and "Art." in df_base.columns:
        df_base = df_base.rename(columns={"Art.": "Artículo"})

    df_base = df_base.dropna(how="all")
    dep = _depurada_operacion_algebra(df_base)

    # Orden final PROCESADO
    col_pol = _find_col_fast(dep, "Póliza", "Poliza")
    col_art = _find_col_fast(dep, "Artículo", "Art.")
    if col_pol and col_art:
        dep = dep.sort_values(by=[col_pol, col_art], ignore_index=True)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    consolidado = out_dir / f"{prefix}_COMPILADO_{ts}.xlsx"
    export_excel_clean(consolidado, {"ORIGINAL": df_orig_clean, "PROCESADO": dep})
    generated_files.append(consolidado)
    log(f"✅ Consolidado creado: {consolidado.name}")

    # Exportar por categoría (filtradas si se eligieron)
    generated_files += _exportar_por_categoria(
        out_dir, prefix, ts, dep, df_orig_clean,
        log=log, prog=prog, cats_filter=cats_filter, prefer_col=prefer_col
    )

    prog(1, 1, "Finalizado")
    return generated_files


# ──────────────────────────────── UI ────────────────────────────────
@dataclass
class AppState:
    files: List[Path]
    out_dir: Path | None
    prefix: str


class ProduccionSoloGUI(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(f"{APP_TITLE} {VERSION}")
        self.minsize(980, 640)
        self.geometry("1060x680")
        try:
            self.tk.call("tk", "scaling", 1.2)  # mejor legibilidad en Windows
        except Exception:
            pass

        self.state = AppState(files=[], out_dir=None, prefix="PROD")

        # cache categorías
        self._prod_cats: List[str] = []
        self._build_ui()

    def _build_ui(self):
        s = ttk.Style(self)
        try:
            s.theme_use("clam")
        except Exception:
            pass

        root = ttk.Frame(self, padding=12)
        root.pack(fill="both", expand=True)
        root.columnconfigure(1, weight=1)
        root.rowconfigure(7, weight=1)

        def row(r, label, widget, btn=None):
            ttk.Label(root, text=label).grid(row=r, column=0, sticky="w", padx=(0, 8), pady=6)
            widget.grid(row=r, column=1, sticky="we", pady=6)
            if btn:
                btn.grid(row=r, column=2, sticky="e", padx=(8, 0), pady=6)

        # Entrada de archivos
        self.var_files = tk.StringVar(value="0 archivo(s) seleccionado(s)")
        row(0, "📁 Archivos de producción (Excel/CSV):",
            ttk.Entry(root, textvariable=self.var_files, state="readonly"),
            ttk.Button(root, text="Agregar…", command=self._add_files))

        # Carpeta salida
        self.var_out = tk.StringVar()
        row(1, "📂 Carpeta de salida:",
            ttk.Entry(root, textvariable=self.var_out, state="readonly"),
            ttk.Button(root, text="Seleccionar…", command=self._choose_out))

        # Prefijo
        self.var_prefix = tk.StringVar(value="PROD")
        row(2, "🏷️ Prefijo archivo(s):", ttk.Entry(root, textvariable=self.var_prefix, width=20))

        # Selector de columna para categorías
        lf_cat = ttk.LabelFrame(root, text="Categorías (Nombre Sección Principal / Nombre Sección)", padding=10)
        lf_cat.grid(row=3, column=0, columnspan=3, sticky="nsew", pady=(6, 6))
        self.var_cat_prefer = tk.StringVar(value="auto")  # auto|principal|seccion

        frm_opt = ttk.Frame(lf_cat)
        frm_opt.pack(fill="x", pady=(0, 6))
        ttk.Label(frm_opt, text="Columna a usar:").pack(side="left")
        ttk.Radiobutton(frm_opt, text="Auto (prioriza Principal)", value="auto",
                        variable=self.var_cat_prefer, command=self._reload_prod_categories_if_any).pack(side="left", padx=6)
        ttk.Radiobutton(frm_opt, text="Nombre Sección Principal", value="principal",
                        variable=self.var_cat_prefer, command=self._reload_prod_categories_if_any).pack(side="left", padx=6)
        ttk.Radiobutton(frm_opt, text="Nombre Sección", value="seccion",
                        variable=self.var_cat_prefer, command=self._reload_prod_categories_if_any).pack(side="left", padx=6)

        frm_btn = ttk.Frame(lf_cat)
        frm_btn.pack(fill="x", pady=(0, 6))
        ttk.Button(frm_btn, text="Cargar categorías", command=self._load_prod_categories).pack(side="left")
        ttk.Button(frm_btn, text="Seleccionar todo", command=lambda: self.list_prod_cats.selection_set(0, tk.END)).pack(side="left", padx=6)
        ttk.Button(frm_btn, text="Limpiar selección", command=lambda: self.list_prod_cats.selection_clear(0, tk.END)).pack(side="left")

        self.list_prod_cats = tk.Listbox(lf_cat, selectmode="extended", height=8, exportselection=False)
        self.list_prod_cats.pack(fill="both", expand=True)

        # Botones ejecutar
        btns = ttk.Frame(root)
        btns.grid(row=4, column=0, columnspan=3, sticky="e", pady=(4, 6))
        self.btn_run = ttk.Button(btns, text="▶ Procesar", command=self._run, width=16)
        self.btn_run.grid(row=0, column=0, padx=6)
        ttk.Button(btns, text="Limpiar lista", command=self._clear_list, width=14).grid(row=0, column=1, padx=6)
        ttk.Button(btns, text="Salir", command=self.destroy, width=10).grid(row=0, column=2, padx=6)

        # Progreso
        self.progress = ttk.Progressbar(root, orient="horizontal", mode="determinate", maximum=100)
        self.progress.grid(row=5, column=0, columnspan=3, sticky="we", pady=(0, 6))

        # Log
        lf = ttk.LabelFrame(root, text="Registro", padding=8)
        lf.grid(row=6, column=0, columnspan=3, sticky="nsew")
        self.log = ScrolledText(lf, height=12, font=("Consolas", 10))
        self.log.pack(fill="both", expand=True)

    # --------------- UI helpers ---------------
    def _add_files(self):
        fl = filedialog.askopenfilenames(title="Seleccionar producción",
                                         filetypes=[("Excel/CSV", "*.xlsx *.xlsm *.xls *.csv")])
        if not fl:
            return
        cur = {p.resolve() for p in self.state.files}
        for p in fl:
            P = Path(p)
            if P.resolve() not in cur:
                self.state.files.append(P)
        self.var_files.set(f"{len(self.state.files)} archivo(s) seleccionado(s)")

    def _choose_out(self):
        d = filedialog.askdirectory(title="Carpeta de salida")
        if d:
            self.state.out_dir = Path(d)
            self.var_out.set(d)

    def _clear_list(self):
        self.state.files = []
        self.var_files.set("0 archivo(s) seleccionado(s)")
        self._prod_cats = []
        self.list_prod_cats.delete(0, tk.END)

    def _log(self, msg: str):
        self.log.insert("end", msg + "\n")
        self.log.see("end")

    def _set_progress(self, val: int):
        self.progress["value"] = max(0, min(100, int(val)))
        self.update_idletasks()

    # --------------- Categorías ---------------
    def _reload_prod_categories_if_any(self):
        if self._prod_cats:
            self._load_prod_categories()

    def _load_prod_categories(self):
        if not self.state.files:
            messagebox.showwarning("Producción", "Agrega archivos de producción primero.")
            return
        dfs: List[pd.DataFrame] = []
        # lee hasta 5 archivos para detectar categorías
        for p in self.state.files[:5]:
            P = Path(p)
            try:
                if P.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
                    df = safe_read_excel(P)
                else:
                    df = safe_read_csv(P)
                if not df.empty:
                    dfs.append(df)
            except Exception:
                continue
        if not dfs:
            messagebox.showwarning("Producción", "No se pudo leer datos para detectar categorías.")
            return
        df_all = pd.concat(dfs, ignore_index=True, sort=False)
        col = _pick_cat_column(df_all, self.var_cat_prefer.get())
        self.list_prod_cats.delete(0, tk.END)
        if not col:
            self._prod_cats = ["(SIN CATEGORÍA – todo)"]
            self.list_prod_cats.insert(tk.END, self._prod_cats[0])
            return
        cats = sorted([c for c in df_all[col].dropna().astype(str).unique()], key=lambda x: str(x))
        self._prod_cats = cats
        for c in cats:
            self.list_prod_cats.insert(tk.END, c)

    def _get_selected_prod_categories(self) -> Optional[Set[str]]:
        if not self._prod_cats:
            return None
        sel = [self.list_prod_cats.get(i) for i in self.list_prod_cats.curselection()]
        if not sel:
            return None  # ninguna seleccionada = todas
        if "(SIN CATEGORÍA" in sel[0]:
            return None
        return set(map(str, sel))

    # --------------- Run ---------------
    def _run(self):
        if not self.state.files:
            messagebox.showerror("Producción", "Agrega uno o más archivos de entrada.")
            return
        if not self.state.out_dir:
            messagebox.showerror("Salida", "Selecciona una carpeta de salida.")
            return
        self.state.prefix = self.var_prefix.get().strip() or "PROD"

        self.btn_run.configure(state="disabled")
        self._set_progress(0)
        self.log.delete("1.0", "end")
        self._log(f"Leyendo {len(self.state.files)} archivo(s)…")

        cats_selected = self._get_selected_prod_categories()
        prefer = self.var_cat_prefer.get()

        threading.Thread(
            target=self._worker,
            args=(cats_selected, prefer),
            daemon=True
        ).start()

    def _worker(self, cats_selected: Optional[Set[str]], prefer_col: str):
        try:
            gen = procesar_produccion(
                self.state.files,
                self.state.out_dir,
                self.state.prefix,
                log=lambda m: self.after(0, self._log, m),
                prog=lambda cur, tot, _: self.after(0, self._set_progress, int((cur / max(tot, 1)) * 100)),
                cats_filter=cats_selected,
                prefer_col=prefer_col
            )
            if gen:
                self.after(0, self._set_progress, 100)
                self.after(0, messagebox.showinfo, "Completado",
                           f"Se generaron {len(gen)} archivo(s) en:\n{self.state.out_dir}")
            else:
                self.after(0, messagebox.showwarning, "Atención", "No se generaron archivos.")
        except Exception as e:
            import traceback
            self.after(0, self._log, f"Error: {e}\n{traceback.format_exc()}")
        finally:
            self.after(0, self.btn_run.configure, {"state": "normal"})


# ──────────────────────────────── Main ────────────────────────────────
if __name__ == "__main__":
    app = ProduccionSoloGUI()
    app.mainloop()
