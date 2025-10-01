#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import re
import sys
import shutil
import threading
import unicodedata
from datetime import datetime, date
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Set

import numpy as np
import pandas as pd

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from tkinter.scrolledtext import ScrolledText

from openpyxl import load_workbook
from openpyxl.styles import NamedStyle
from openpyxl.descriptors import String as _StringDesc

APP_TITLE = "Tarifas – SOLO Producción (con selector)"
VERSION   = "v1.2"
_StringDesc.allow_none = True

# ───────── util texto ─────────
def _norm(s: str) -> str:
    if s is None:
        return ""
    s = "".join(c for c in unicodedata.normalize("NFKD", str(s)) if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()

def normalize_token(x) -> str:
    """Normaliza cabeceras: quita acentos, TODO espacio/salto de línea y signos; sólo [a-z0-9]."""
    if x is None:
        return ""
    s = "".join(c for c in unicodedata.normalize("NFKD", str(x)) if not unicodedata.combining(c))
    s = s.lower()
    s = re.sub(r"\s+", "", s, flags=re.UNICODE)   # elimina espacios, tabs y saltos de línea
    s = re.sub(r"[^a-z0-9]", "", s)               # elimina signos (/, ., :, etc.)
    return s

def _find_col_fast(df: pd.DataFrame, *cands: str) -> Optional[str]:
    if df is None or df.empty:
        return None
    norm_map = {_norm(c): c for c in df.columns}
    for cand in cands:
        k = _norm(cand)
        if k in norm_map:
            return norm_map[k]
    for cand in cands:
        k = _norm(cand)
        for kk, real in norm_map.items():
            if k and k in kk:
                return real
    return None

# ───────── lectura robusta (lo necesario) ─────────
def open_and_clean_workbook(path: Path):
    wb = load_workbook(str(path), read_only=True, data_only=True, rich_text=False, keep_links=False)
    valid = [st for st in getattr(wb, "_named_styles", []) if isinstance(getattr(st, "name", None), str)]
    if "Normal" not in [st.name for st in valid]:
        valid.insert(0, NamedStyle(name="Normal"))
    wb._named_styles = valid
    wb._named_style_map = {st.name: i for i, st in enumerate(valid)}
    if hasattr(wb, "named_styles"):
        wb.named_styles = valid
    return wb

def _scan_for_header_row(values_2d, max_rows=80) -> int:
    keywords = {
        "poliza","póliza","endoso","art","artículo","articulo",
        "f/emision","f emision","f emisión","fecha emisión",
        "fec. desde art","fec. hasta art",
        "nombre sección","nombre seccion","sección principal","seccion principal"
    }
    best_row, best_score = 0, -1
    upto = min(max_rows, len(values_2d))
    for r in range(upto):
        row = values_2d[r]
        toks = [_norm(x) for x in row if x not in (None, "")]
        if not toks:
            continue
        kw_hits = sum(any(k in t for k in keywords) for t in toks)
        score = kw_hits * 3 + len(set(toks))
        if score > best_score:
            best_score, best_row = score, r
    return best_row

def _sheet_to_dataframe(ws) -> pd.DataFrame:
    values = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    if not values:
        return pd.DataFrame()
    hdr_idx = _scan_for_header_row(values, max_rows=80)
    header  = values[hdr_idx]
    data    = [r for r in values[hdr_idx+1:] if any(x not in (None,""," ") for x in r)]
    cols    = [str(c).strip() if c is not None else f"Col_{i}" for i, c in enumerate(header)]
    return pd.DataFrame(data, columns=cols)

def robust_read_excel_all_sheets(path: Path) -> pd.DataFrame:
    try:
        wb = open_and_clean_workbook(path)
        frames = []
        for sh in wb.sheetnames:
            try:
                df = _sheet_to_dataframe(wb[sh])
                if not df.empty:
                    frames.append(df)
            except Exception:
                continue
        wb.close()
        if frames:
            return pd.concat(frames, ignore_index=True, sort=False)
    except Exception:
        pass
    try:
        xls = pd.ExcelFile(path, engine="openpyxl")
        frames = []
        for sh in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sh, header=0)
            if not df.empty:
                frames.append(df)
        if frames:
            return pd.concat(frames, ignore_index=True, sort=False)
    except Exception:
        pass
    return pd.DataFrame()

def safe_read_excel_prefer(path: Path, prefer_sheets: List[str] | None = None) -> pd.DataFrame:
    try:
        wb = open_and_clean_workbook(path)
        ws = None
        if prefer_sheets:
            low = {n.lower(): n for n in wb.sheetnames}
            for want in prefer_sheets:
                if want.lower() in low:
                    ws = wb[low[want.lower()]]
                    break
        if ws is None:
            ws = wb.active
        df = _sheet_to_dataframe(ws)
        wb.close()
        if not df.empty:
            return df
        return robust_read_excel_all_sheets(path)
    except Exception:
        return robust_read_excel_all_sheets(path)

# ───────── pegado por cabeceras ─────────
def _build_synonyms_map_prod() -> Dict[str, List[str]]:
    # Claves = nombres CANÓNICOS tal como aparecen (idealmente) en la base
    return {
        "Póliza": ["Póliza","Poliza"],
        "Endoso": ["Endoso"],
        # Variantes amplias para cubrir /, acentos, con/sin “de” y saltos de línea
        "Fecha de Emisión": [
            "Fecha de Emisión","Fecha de Emision","Fecha Emisión","Fecha Emision",
            "F/Emisión","F/Emision","F Emisión","F Emision","f emision","f/emision"
        ],
        "Nombre Agente Principal": ["Nombre Agente Principal","Agente","Agente Principal"],
        "Artículo": ["Artículo","Art.","Articulo"],
        "Fec. Desde Art.": ["Fec. Desde Art.","Fecha Desde Art.","Inicio Vigencia"],
        "Fec. Hasta Art.": ["Fec. Hasta Art.","Fecha Hasta Art.","Fin Vigencia"],
        "Prima Art.": ["Prima Art.","Prima artículo","Prima Articulo"],
        "Suma Asegurada Art.": ["Suma Asegurada Art.","Suma asegurada art.","Suma Asegurada Articulo"],
        "fechaInicio": ["fechaInicio","Fecha Inicio","Fec. Desde Art.","Fecha Desde Art."],
        "fechaFin": ["fechaFin","Fecha Fin","Fec. Hasta Art.","Fecha Hasta Art."],
        "Suma Art.": ["Suma Art.","Suma Prima Art.","Prima Total Art.","Suma de Prima Art."],
        "Suma Costo de Servicio Art.": ["Suma Costo de Servicio Art.","Suma Costo Servicio","Costo de Servicio Art."],
        "Suma Importe Agente Art.": ["Suma Importe Agente Art.","Suma Importe Agente","Importe Agente Art."],
    }

DATE_CANONICAL = {
    "Fecha de Emisión","Fec. Desde Art.","Fec. Hasta Art.","fechaInicio","fechaFin"
}

def _normalize_dict_keys(d: Dict[str, List[str]]) -> Dict[str, List[str]]:
    return {k: [normalize_token(x) for x in (v + [k])] for k, v in d.items()}

def _find_header_row_and_map(ws, synonyms: Dict[str, List[str]]) -> Tuple[int, Dict[str, int]]:
    """
    Busca la fila de encabezados y devuelve:
      (fila_header_1based, mapeo {Nombre Canónico -> col_index_1based})
    """
    syn_norm = _normalize_dict_keys(synonyms)
    best_row, best_hits, best_map = None, -1, {}
    for r in range(1, min(ws.max_row, 80) + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        tokens = [normalize_token(v) for v in vals]
        local, hits = {}, 0
        for c_idx, tok in enumerate(tokens, start=1):
            if not tok:
                continue
            for canon, opts in syn_norm.items():
                if tok in opts and canon not in local:
                    local[canon] = c_idx
                    hits += 1
        if hits > best_hits and hits >= 3:
            best_row, best_hits, best_map = r, hits, local
    if best_row is None:
        best_row = 1
        vals = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        tokens = [normalize_token(v) for v in vals]
        best_map = {}
        for c_idx, tok in enumerate(tokens, start=1):
            for canon, opts in syn_norm.items():
                if tok in opts and canon not in best_map:
                    best_map[canon] = c_idx
    return best_row, best_map

def _clear_below(ws, header_row: int, col_indices: List[int]):
    """Opcional: limpia debajo del encabezado sólo las columnas a escribir."""
    max_row = ws.max_row
    for r in range(header_row + 1, max_row + 1):
        for c in col_indices:
            ws.cell(row=r, column=c).value = None

def _maybe_parse_date(val):
    if val is None or (isinstance(val, float) and np.isnan(val)):
        return None
    if isinstance(val, (datetime, date)):
        return val
    s = str(val).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            pass
    # intento flexible
    try:
        return pd.to_datetime(s, dayfirst=True, errors="coerce").date()
    except Exception:
        return val

def _paste_df_to_sheet(ws, df: pd.DataFrame, synonyms: Dict[str, List[str]], log) -> int:
    header_row, canon2col = _find_header_row_and_map(ws, synonyms)
    if not canon2col:
        log("⚠️ No se detectaron encabezados en la hoja destino ('Prod').")
        return 0

    syn_norm = _normalize_dict_keys(synonyms)
    df_norm_map = {normalize_token(c): c for c in df.columns}

    # df_col_name -> sheet_col_index
    final_map: List[Tuple[str, int, str]] = []
    for canon, c_idx in canon2col.items():
        aliases = syn_norm.get(canon, [])
        hit = None
        for alias in aliases:
            if alias in df_norm_map:
                hit = df_norm_map[alias]
                break
        if hit is None and normalize_token(canon) in df_norm_map:
            hit = df_norm_map[normalize_token(canon)]
        if hit is not None:
            final_map.append((hit, c_idx, canon))

    if not final_map:
        log("⚠️ No hay columnas coincidentes entre los datos y la hoja destino.")
        return 0

    # Limpieza previa sólo en columnas a escribir (descomenta si quieres limpiar)
    # _clear_below(ws, header_row, [c for _, c, _ in final_map])

    # Pegar comenzando justo debajo del encabezado
    count = 0
    records = df.replace({np.nan: None}).to_dict("records")
    for i, rec in enumerate(records, start=1):
        r = header_row + i
        for df_col, c_idx, canon in final_map:
            val = rec.get(df_col)
            if canon in DATE_CANONICAL:
                val = _maybe_parse_date(val)
            cell = ws.cell(row=r, column=c_idx, value=val)
            if isinstance(val, (date, datetime)):
                cell.number_format = "DD/MM/YYYY"
        count += 1
    return count

# ───────── núcleo: generar tarifas desde producción con filtro de categorías ─────────
def _pick_cat_column(df: pd.DataFrame, prefer: str = "auto") -> Optional[str]:
    if prefer == "principal":
        return _find_col_fast(df, "Nombre Sección Principal", "Nombre Seccion Principal")
    if prefer == "seccion":
        return _find_col_fast(df, "Nombre Sección", "Nombre Seccion")
    return (_find_col_fast(df, "Nombre Sección Principal", "Nombre Seccion Principal")
            or _find_col_fast(df, "Nombre Sección", "Nombre Seccion"))

def generar_tarifas_desde_produccion(
    prod_compilado: Path,
    base_monocobertura: Path,
    out_dir: Path,
    prefix: str,
    log,
    prog,
    categorias_seleccionadas: Optional[Set[str]] = None,
    prefer_col: str = "auto",
) -> List[Path]:
    out: List[Path] = []
    log("Cargando producción (PROCESADO)…")
    df_prod = safe_read_excel_prefer(prod_compilado, ["PROCESADO", "Prod_procesados", "PROD_PROCESADOS"])
    if df_prod.empty:
        log("⚠️ No se encontró hoja PROCESADO ni datos válidos en producción.")
        return out

    cat_col = _pick_cat_column(df_prod, prefer_col)
    if cat_col:
        cats_all = sorted([c for c in pd.Series(df_prod[cat_col]).dropna().astype(str).unique()], key=lambda x: str(x))
        cats = cats_all
        if categorias_seleccionadas:
            target = {str(x) for x in categorias_seleccionadas}
            cats = [c for c in cats_all if str(c) in target]
            if not cats:
                log("⚠️ Ninguna de las categorías seleccionadas está presente en el archivo. Se cancela.")
                return out
    else:
        cats = ["TODAS"]
        df_prod["_CAT_"] = "TODAS"
        cat_col = "_CAT_"

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    total = len(cats)
    log(f"Categorías: {total}")

    for i, cat in enumerate(cats, 1):
        prog(i, total, f"Pegando {cat}")
        df_cat = df_prod[df_prod[cat_col].astype(str) == str(cat)] if cat_col else df_prod.copy()
        if df_cat.empty:
            log(f"  · {cat}: sin filas; se omite.")
            continue

        safe = re.sub(r'[\\/*?:"<>|]', "_", str(cat)).strip()[:60] or "CAT"
        out_path = out_dir / f"{prefix}_{safe}_{ts}.xlsx"
        shutil.copy(base_monocobertura, out_path)

        try:
            wb = load_workbook(out_path, data_only=False, keep_links=True)
            if "Prod" not in wb.sheetnames:
                log("⚠️ El libro base no tiene hoja 'Prod'.")
                wb.close()
                continue
            ws = wb["Prod"]
            pasted = _paste_df_to_sheet(ws, df_cat, _build_synonyms_map_prod(), log)
            wb.save(out_path)
            wb.close()
            out.append(out_path)
            log(f"  · {cat}: {pasted} fila(s) pegadas → {out_path.name}")
        except Exception as e:
            log(f"  · {cat}: error escribiendo '{out_path.name}': {e}")

    prog(total, total, "Finalizado")
    return out

# ───────── UI ─────────
class SoloTarifasGUI(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(f"{APP_TITLE} {VERSION}")
        self.geometry("1060x640")
        self.minsize(980, 600)

        self.var_prod   = tk.StringVar()
        self.var_base   = tk.StringVar()
        self.var_dest   = tk.StringVar()
        self.var_prefix = tk.StringVar(value="TARIFAS")

        # Selector categorías
        self.var_prefer_col = tk.StringVar(value="auto")   # auto | principal | seccion
        self.list_cats = None  # Listbox
        self._cats: List[str] = []  # cache categorías
        self._build_ui()

    def _build_ui(self):
        root = ttk.Frame(self, padding=12); root.pack(fill="both", expand=True)
        root.columnconfigure(1, weight=1)

        def row(r, label, widget, btn=None):
            ttk.Label(root, text=label).grid(row=r, column=0, sticky="w", padx=(0,8), pady=6)
            widget.grid(row=r, column=1, sticky="we", pady=6)
            if btn: btn.grid(row=r, column=2, padx=(8,0))

        row(0, "📄 Producción PROCESADA (.xlsx):",
            ttk.Entry(root, textvariable=self.var_prod),
            ttk.Button(root, text="Examinar", command=lambda: self._pick(self.var_prod, "*.xlsx")))

        row(1, "📘 Base monocobertura (.xlsx):",
            ttk.Entry(root, textvariable=self.var_base),
            ttk.Button(root, text="Examinar", command=lambda: self._pick(self.var_base, "*.xlsx")))

        row(2, "📂 Carpeta de salida:",
            ttk.Entry(root, textvariable=self.var_dest),
            ttk.Button(root, text="Seleccionar", command=self._pick_dir))

        row(3, "🏷️ Prefijo archivos:",
            ttk.Entry(root, textvariable=self.var_prefix, width=20))

        # ── Selector de categorías
        lf = ttk.LabelFrame(root, text="Categorías (desde PROCESADO)", padding=10)
        lf.grid(row=4, column=0, columnspan=3, sticky="nsew", pady=(6,6))
        root.rowconfigure(4, weight=1)

        frm_opt = ttk.Frame(lf); frm_opt.pack(fill="x", pady=(0,6))
        ttk.Label(frm_opt, text="Columna a usar:").pack(side="left")
        for val, txt in (("auto","Auto (prioriza Principal)"),
                         ("principal","Nombre Sección Principal"),
                         ("seccion","Nombre Sección")):
            ttk.Radiobutton(frm_opt, text=txt, value=val, variable=self.var_prefer_col,
                            command=self._reload_categories_if_any).pack(side="left", padx=6)

        frm_btn = ttk.Frame(lf); frm_btn.pack(fill="x", pady=(0,6))
        ttk.Button(frm_btn, text="Cargar categorías", command=self._load_categories).pack(side="left")
        ttk.Button(frm_btn, text="Seleccionar todo", command=self._select_all).pack(side="left", padx=6)
        ttk.Button(frm_btn, text="Limpiar selección", command=self._clear_sel).pack(side="left")

        self.list_cats = tk.Listbox(lf, selectmode="extended", height=10, exportselection=False)
        self.list_cats.pack(fill="both", expand=True)

        btns = ttk.Frame(root); btns.grid(row=5, column=0, columnspan=3, sticky="e", pady=(8,6))
        self.btn_run = ttk.Button(btns, text="Generar libros", command=self._run, width=18)
        self.btn_run.grid(row=0, column=0, padx=6)
        ttk.Button(btns, text="Salir", command=self.destroy, width=10).grid(row=0, column=1, padx=6)

        self.progress = ttk.Progressbar(root, orient="horizontal", mode="determinate", maximum=100)
        self.progress.grid(row=6, column=0, columnspan=3, sticky="we", pady=(0,8))

        lf_log = ttk.LabelFrame(root, text="Registro", padding=8)
        lf_log.grid(row=7, column=0, columnspan=3, sticky="nsew")
        root.rowconfigure(7, weight=1)
        self.log = ScrolledText(lf_log, height=10, font=("Consolas", 10))
        self.log.pack(fill="both", expand=True)

    def _pick(self, var: tk.StringVar, pattern="*.*"):
        f = filedialog.askopenfilename(filetypes=[("Archivo", pattern)])
        if f: var.set(f)

    def _pick_dir(self):
        d = filedialog.askdirectory()
        if d: self.var_dest.set(d)

    def _log(self, msg: str):
        self.log.insert("end", msg + "\n"); self.log.see("end")

    def _set_progress(self, val: int):
        self.progress["value"] = max(0, min(100, int(val))); self.update_idletasks()

    # ── Categorías
    def _load_categories(self):
        prod = self.var_prod.get().strip()
        if not prod:
            messagebox.showwarning("Producción", "Selecciona el archivo de producción PROCESADA.")
            return
        df = safe_read_excel_prefer(Path(prod), ["PROCESADO", "Prod_procesados", "PROD_PROCESADOS"])
        if df.empty:
            messagebox.showwarning("Producción", "No se pudo leer PROCESADO.")
            return
        col = _pick_cat_column(df, self.var_prefer_col.get())
        if not col:
            self._cats = ["(SIN CATEGORÍA – todo)"]
            self._refresh_listbox(self._cats)
            return
        cats = sorted([c for c in pd.Series(df[col]).dropna().astype(str).unique()], key=lambda x: str(x))
        self._cats = cats
        self._refresh_listbox(cats)

    def _reload_categories_if_any(self):
        if self._cats:
            self._load_categories()

    def _refresh_listbox(self, cats: List[str]):
        self.list_cats.delete(0, tk.END)
        for c in cats:
            self.list_cats.insert(tk.END, c)

    def _select_all(self):
        self.list_cats.selection_set(0, tk.END)

    def _clear_sel(self):
        self.list_cats.selection_clear(0, tk.END)

    def _get_selected_categories(self) -> Optional[Set[str]]:
        if not self._cats:
            return None
        sel = [self.list_cats.get(i) for i in self.list_cats.curselection()]
        if not sel:
            return None  # ninguna seleccionada = todas
        if "(SIN CATEGORÍA" in sel[0]:
            return None
        return set(map(str, sel))

    # ── Run
    def _run(self):
        prod = Path(self.var_prod.get().strip())
        base = Path(self.var_base.get().strip())
        dest = Path(self.var_dest.get().strip()) if self.var_dest.get().strip() else None
        if not prod.exists():
            messagebox.showerror("Producción", "Selecciona el archivo de producción PROCESADA (.xlsx)."); return
        if not base.exists():
            messagebox.showerror("Base", "Selecciona el libro base monocobertura (.xlsx)."); return
        if dest is None:
            messagebox.showerror("Salida", "Selecciona una carpeta de salida."); return
        dest.mkdir(parents=True, exist_ok=True)

        self.btn_run.configure(state="disabled")
        self._set_progress(0)
        self.log.delete("1.0", "end")
        self._log("Iniciando…")

        sel = self._get_selected_categories()
        prefer = self.var_prefer_col.get()

        threading.Thread(
            target=self._worker,
            args=(prod, base, dest, self.var_prefix.get().strip() or "TARIFAS", sel, prefer),
            daemon=True
        ).start()

    def _worker(self, prod, base, dest, pref, sel, prefer):
        try:
            generados = generar_tarifas_desde_produccion(
                prod, base, dest, pref,
                log=lambda m: self.after(0, self._log, m),
                prog=lambda cur, tot, _: self.after(0, self._set_progress, int((cur/max(tot,1))*100)),
                categorias_seleccionadas=sel,
                prefer_col=prefer
            )
            if generados:
                self.after(0, self._set_progress, 100)
                self.after(0, messagebox.showinfo, "Completado",
                           f"Se generaron {len(generados)} archivo(s) en:\n{dest}")
            else:
                self.after(0, messagebox.showwarning, "Atención", "No se generaron libros.")
        except Exception as e:
            import traceback
            self.after(0, self._log, f"Error: {e}\n{traceback.format_exc()}")
        finally:
            self.after(0, self.btn_run.configure, {"state": "normal"})

if __name__ == "__main__":
    app = SoloTarifasGUI()
    app.mainloop()
